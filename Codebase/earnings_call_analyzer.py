#!/usr/bin/env python3
"""
Earnings Call Analyzer for Alpha OS - Terminal Edition
Analyzes earnings call transcripts and populates Excel with comprehensive analysis

Usage:
    export ANTHROPIC_API_KEY="your-key-here"
    python earnings_call_analyzer.py input.xlsx [output.xlsx]
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import requests
import json
import os
from datetime import datetime
import time
import re
import sys

def get_api_key():
    """Get Anthropic API key from environment"""
    key = os.getenv('ANTHROPIC_API_KEY')
    if not key:
        print("\n❌ ERROR: ANTHROPIC_API_KEY not set")
        print("\nSet it with:")
        print('  export ANTHROPIC_API_KEY="sk-ant-..."')
        print("\nOr run with:")
        print('  ANTHROPIC_API_KEY="sk-ant-..." python earnings_call_analyzer.py input.xlsx\n')
        sys.exit(1)
    return key

def search_and_analyze_earnings_call(ticker, company_name, api_key, max_retries=2):
    """
    Use Claude API with web_search tool to find and analyze earnings calls
    Returns dict with analysis or None if failed
    """
    
    for attempt in range(max_retries):
        try:
            prompt = f"""Find and analyze the most recent earnings call transcript for {company_name} ({ticker}).

Steps:
1. Search for recent earnings call transcript (Seeking Alpha, Motley Fool, company IR pages)
2. Read the full transcript
3. Analyze comprehensively

Return analysis in this JSON format (keep field lengths under limits):

{{
  "transcript_url": "URL of transcript source",
  "transcript_date": "Date of earnings call (e.g., 'Q4 2024', 'Feb 2025')",
  "forward_guidance": "Management outlook, revenue/earnings guidance for next quarters (max 200 chars)",
  "strategic_initiatives": "Key projects, M&A, new products, strategic shifts mentioned (max 150 chars)",
  "analyst_questions_themes": "Main analyst concerns and question topics (max 150 chars)",
  "management_response_tone": "Confident/Defensive/Evasive/Optimistic/Cautious with brief reason (max 100 chars)",
  "beat_miss_vs_expectations": "Beat/miss/met on revenue and EPS vs consensus (max 100 chars)",
  "margin_profitability": "Gross margin, operating margin, profitability trends mentioned (max 100 chars)",
  "capex_investments": "Capital expenditure plans, investment priorities (max 100 chars)",
  "competitive_positioning": "How they position vs competitors, market share (max 100 chars)",
  "risk_flags": "Concerns, headwinds, risks mentioned (max 200 chars)",
  "key_takeaways": "3-5 most important bullet points (max 400 chars total)",
  "overall_sentiment": "Bullish/Neutral/Bearish with 1-sentence justification (max 150 chars)"
}}

If no transcript found, return: {{"transcript_url": "Not Found", "forward_guidance": "No recent earnings call available"}}

Return ONLY valid JSON, no preamble or markdown."""

            response = requests.post(
                'https://api.anthropic.com/v1/messages',
                headers={
                    'x-api-key': api_key,
                    'anthropic-version': '2023-06-01',
                    'content-type': 'application/json'
                },
                json={
                    'model': 'claude-sonnet-4-20250514',  # Sonnet 4 - your API has access
                    'max_tokens': 4000,
                    'tools': [
                        {
                            'type': 'web_search_20250305',
                            'name': 'web_search'
                        }
                    ],
                    'messages': [{'role': 'user', 'content': prompt}]
                },
                timeout=180
            )
            
            # Handle rate limit errors
            if response.status_code == 429:
                if attempt < max_retries - 1:
                    wait_time = 45  # Fixed 45 second wait on rate limit
                    print(f"    ⚠️  Rate limited, waiting {wait_time}s before retry...")
                    time.sleep(wait_time)
                    continue
                else:
                    print(f"    ❌ Rate limit - skipping after {max_retries} attempts")
                    return None
            
            if response.status_code != 200:
                print(f"    ❌ API error {response.status_code}: {response.text[:200]}")
                return None
            
            result = response.json()
            
            # Extract text from all content blocks
            full_text = ""
            for block in result.get('content', []):
                if block.get('type') == 'text':
                    full_text += block.get('text', '')
            
            if not full_text:
                print(f"    ❌ No text in response")
                return None
            
            # Extract JSON (handle markdown code blocks)
            clean_text = full_text.strip()
            if clean_text.startswith('```'):
                # Remove markdown code block
                clean_text = re.sub(r'^```(?:json)?\s*', '', clean_text)
                clean_text = re.sub(r'```\s*$', '', clean_text)
            
            # Find JSON object
            json_match = re.search(r'\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}', clean_text, re.DOTALL)
            if not json_match:
                print(f"    ❌ No JSON found in response")
                return None
            
            analysis = json.loads(json_match.group())
            return analysis
            
        except json.JSONDecodeError as e:
            print(f"    ❌ JSON parse error: {str(e)}")
            if attempt < max_retries - 1:
                time.sleep(5)
                continue
            return None
        except requests.exceptions.Timeout:
            print(f"    ❌ Request timeout")
            if attempt < max_retries - 1:
                time.sleep(10)
                continue
            return None
        except Exception as e:
            print(f"    ❌ Error: {type(e).__name__}: {str(e)}")
            if attempt < max_retries - 1:
                time.sleep(5)
                continue
            return None
    
    return None  # All retries exhausted

def process_screening_excel(input_file, output_file, api_key):
    """Read screening Excel and add earnings call analysis"""
    
    print("\n" + "="*80)
    print("EARNINGS CALL ANALYZER")
    print("="*80)
    print(f"Input:  {input_file}")
    print(f"Output: {output_file}")
    print("="*80 + "\n")
    
    # Read Excel - handle multi-sheet files
    try:
        excel_file = pd.ExcelFile(input_file)
        sheet_names = excel_file.sheet_names
        
        print(f"📑 Found {len(sheet_names)} sheets: {', '.join(sheet_names)}\n")
        
        # If multiple sheets, combine all non-summary sheets
        if len(sheet_names) > 1:
            all_dfs = []
            for sheet in sheet_names:
                if sheet.lower() == 'summary':
                    continue
                df_sheet = pd.read_excel(input_file, sheet_name=sheet)
                # Add sheet name as source
                df_sheet['Source_Sheet'] = sheet
                all_dfs.append(df_sheet)
            
            if not all_dfs:
                print("❌ No data sheets found (only Summary)\n")
                sys.exit(1)
            
            df = pd.concat(all_dfs, ignore_index=True)
            print(f"✅ Combined {len(all_dfs)} sector sheets\n")
        else:
            df = pd.read_excel(input_file)
        
    except Exception as e:
        print(f"❌ Failed to read {input_file}: {e}\n")
        sys.exit(1)
    
    print(f"📊 Found {len(df)} companies to analyze\n")
    
    # Find ticker column (try multiple variations)
    ticker_col = None
    for col in ['Ticker', 'ticker', 'TICKER', 'Symbol', 'symbol']:
        if col in df.columns:
            ticker_col = col
            break
    
    if not ticker_col:
        print(f"❌ Excel must have a ticker/symbol column")
        print(f"   Found columns: {df.columns.tolist()}\n")
        sys.exit(1)
    
    print(f"✅ Using '{ticker_col}' column for tickers\n")
    
    # Add earnings analysis columns
    new_cols = [
        'EC_Transcript_URL',
        'EC_Date', 
        'EC_Forward_Guidance',
        'EC_Strategic_Initiatives',
        'EC_Analyst_Questions',
        'EC_Mgmt_Response_Tone',
        'EC_Beat_Miss',
        'EC_Margin_Commentary',
        'EC_Capex_Plans',
        'EC_Competitive_Position',
        'EC_Risk_Flags',
        'EC_Key_Takeaways',
        'EC_Overall_Sentiment',
        'EC_Analysis_Timestamp'
    ]
    
    for col in new_cols:
        if col not in df.columns:
            df[col] = ''
    
    # Process each company
    success_count = 0
    not_found_count = 0
    failed_count = 0
    
    # Find company name column
    company_col = None
    for col in ['Company Name', 'company_name', 'Company', 'company', 'Name', 'name']:
        if col in df.columns:
            company_col = col
            break
    
    for idx, row in df.iterrows():
        ticker = str(row.get(ticker_col, '')).strip().upper()
        company_name = row.get(company_col, '') if company_col else ticker
        
        if not ticker:
            continue
        
        print(f"[{idx+1}/{len(df)}] {ticker} - {company_name}")
        
        # Run analysis
        analysis = search_and_analyze_earnings_call(ticker, company_name, api_key)
        
        timestamp = datetime.now().strftime('%Y-%m-%d %H:%M')
        
        if not analysis:
            print(f"    ❌ Analysis failed\n")
            df.at[idx, 'EC_Transcript_URL'] = 'Analysis Failed'
            df.at[idx, 'EC_Analysis_Timestamp'] = timestamp
            failed_count += 1
            continue
        
        # Check if transcript found
        if analysis.get('transcript_url') == 'Not Found':
            print(f"    ⚠️  No transcript found\n")
            df.at[idx, 'EC_Transcript_URL'] = 'Not Found'
            df.at[idx, 'EC_Forward_Guidance'] = analysis.get('forward_guidance', 'No transcript')
            df.at[idx, 'EC_Analysis_Timestamp'] = timestamp
            not_found_count += 1
            continue
        
        # Populate all fields
        df.at[idx, 'EC_Transcript_URL'] = analysis.get('transcript_url', '')[:500]
        df.at[idx, 'EC_Date'] = analysis.get('transcript_date', '')
        df.at[idx, 'EC_Forward_Guidance'] = analysis.get('forward_guidance', '')[:1000]
        df.at[idx, 'EC_Strategic_Initiatives'] = analysis.get('strategic_initiatives', '')[:500]
        df.at[idx, 'EC_Analyst_Questions'] = analysis.get('analyst_questions_themes', '')[:500]
        df.at[idx, 'EC_Mgmt_Response_Tone'] = analysis.get('management_response_tone', '')[:200]
        df.at[idx, 'EC_Beat_Miss'] = analysis.get('beat_miss_vs_expectations', '')[:300]
        df.at[idx, 'EC_Margin_Commentary'] = analysis.get('margin_profitability', '')[:300]
        df.at[idx, 'EC_Capex_Plans'] = analysis.get('capex_investments', '')[:300]
        df.at[idx, 'EC_Competitive_Position'] = analysis.get('competitive_positioning', '')[:300]
        df.at[idx, 'EC_Risk_Flags'] = analysis.get('risk_flags', '')[:500]
        df.at[idx, 'EC_Key_Takeaways'] = analysis.get('key_takeaways', '')[:800]
        df.at[idx, 'EC_Overall_Sentiment'] = analysis.get('overall_sentiment', '')[:200]
        df.at[idx, 'EC_Analysis_Timestamp'] = timestamp
        
        sentiment = analysis.get('overall_sentiment', 'N/A')[:60]
        print(f"    ✅ {sentiment}")
        print(f"       {analysis.get('transcript_url', '')[:65]}...\n")
        
        success_count += 1
        
        # Rate limit - 20 second delay for Sonnet 4 (safest)
        time.sleep(20)
    
    # Save with formatting
    print("="*80)
    print("Saving results...")
    
    try:
        # If original file had multiple sheets, preserve them and add analysis sheet
        if len(sheet_names) > 1:
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                # Copy original sheets
                for sheet in sheet_names:
                    original_df = pd.read_excel(input_file, sheet_name=sheet)
                    original_df.to_excel(writer, sheet_name=sheet, index=False)
                
                # Add new earnings analysis sheet
                df.to_excel(writer, index=False, sheet_name='Earnings Analysis')
                
                wb = writer.book
                ws = writer.sheets['Earnings Analysis']
                
                # Bloomberg terminal style header
                header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
                header_font = Font(bold=True, color='FFFFFF', size=10, name='Arial')
                
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
                # Column widths (adjust based on number of columns)
                widths = {
                    'A': 10, 'B': 28, 'C': 18, 'D': 15, 'E': 12, 'F': 12,
                    'G': 45, 'H': 12, 'I': 45, 'J': 35, 'K': 35, 'L': 20,
                    'M': 25, 'N': 25, 'O': 25, 'P': 25, 'Q': 35, 'R': 50,
                    'S': 30, 'T': 16, 'U': 12, 'V': 12, 'W': 12, 'X': 12,
                    'Y': 12, 'Z': 12
                }
                
                for col, width in widths.items():
                    ws.column_dimensions[col].width = width
                
                # Format data rows
                data_font = Font(size=9, name='Arial')
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(vertical='top', wrap_text=True)
                        cell.font = data_font
                
                # Freeze panes
                ws.freeze_panes = 'A2'
        else:
            # Single sheet file
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Earnings Analysis')
                
                wb = writer.book
                ws = writer.sheets['Earnings Analysis']
                
                # Bloomberg terminal style header
                header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
                header_font = Font(bold=True, color='FFFFFF', size=10, name='Arial')
                
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
                # Column widths
                widths = {
                    'A': 10, 'B': 28, 'C': 18, 'D': 15, 'E': 12, 'F': 12,
                    'G': 45, 'H': 12, 'I': 45, 'J': 35, 'K': 35, 'L': 20,
                    'M': 25, 'N': 25, 'O': 25, 'P': 25, 'Q': 35, 'R': 50,
                    'S': 30, 'T': 16
                }
                
                for col, width in widths.items():
                    ws.column_dimensions[col].width = width
                
                # Format data rows
                data_font = Font(size=9, name='Arial')
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(vertical='top', wrap_text=True)
                        cell.font = data_font
                
                # Freeze panes
                ws.freeze_panes = 'A2'
        
        print(f"✅ Saved: {output_file}")
        
    except Exception as e:
        print(f"❌ Save failed: {e}\n")
        sys.exit(1)
    
    # Summary
    print("\n" + "="*80)
    print("SUMMARY")
    print("="*80)
    print(f"Total companies:       {len(df)}")
    print(f"✅ Analyzed:           {success_count}")
    print(f"⚠️  No transcript:      {not_found_count}")
    print(f"❌ Failed:             {failed_count}")
    print("="*80 + "\n")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("\nUsage:")
        print("  python earnings_call_analyzer.py <input.xlsx> [output.xlsx]")
        print("\nExample:")
        print("  export ANTHROPIC_API_KEY='sk-ant-...'")
        print("  python earnings_call_analyzer.py screening_results.xlsx\n")
        sys.exit(1)
    
    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else input_file.replace('.xlsx', '_with_earnings.xlsx')
    
    if not os.path.exists(input_file):
        print(f"\n❌ File not found: {input_file}\n")
        sys.exit(1)
    
    api_key = get_api_key()
    process_screening_excel(input_file, output_file, api_key)
